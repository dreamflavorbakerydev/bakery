"""Explicit, atomic master-data imports. No fuzzy matching or inferred mappings."""
import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from bakery import models as m
from bakery.services.audit import record_change, snapshot
from bakery.services.catalog import product_is_available


class ImportFailure(Exception):
    pass


@dataclass
class ImportReport:
    inserted: int = 0
    updated: int = 0
    skipped: int = 0


def boolean(value):
    if isinstance(value, bool):
        return value
    if str(value).lower() in ('true', '1', 'yes'):
        return True
    if str(value).lower() in ('false', '0', 'no'):
        return False
    raise ValueError('Expected true/false, yes/no, or 1/0.')


def integer(value):
    result = Decimal(str(value))
    if not result.is_finite() or result != result.to_integral_value():
        raise ValueError('Expected a whole number.')
    return int(result)


def money(value):
    result = Decimal(str(value))
    if not result.is_finite() or result != result.quantize(Decimal('.01')):
        raise ValueError('Expected a finite price with at most two decimal places.')
    return result.quantize(Decimal('.01'))


def calendar_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def clock_time(value):
    return value if isinstance(value, time) else time.fromisoformat(str(value))


def weekdays(value):
    return [integer(part.strip()) for part in str(value).split(',')]


# Required columns, optional columns. Blank optional cells leave existing values intact.
SCHEMAS = {
    'customers': (m.Customer, ['customer_name', 'pricing_zone'],
        ['customer_code', 'customer_id', 'address', 'city', 'state', 'zip', 'phone', 'active', 'timezone', 'payment_terms', 'notes']),
    'locations': (m.CustomerLocation, ['location_code', 'location_name'],
        ['customer_code', 'customer_id', 'store_number', 'address', 'city', 'state', 'zip', 'phone', 'active']),
    'products': (m.Product, ['internal_product_code', 'product_name', 'pack_size', 'unit_of_measure'],
        ['description', 'upc', 'sku', 'size', 'units_per_tray', 'category', 'production_lead_days', 'active', 'sort_order']),
    'customer_products': (m.CustomerProduct, ['internal_product_code'],
        ['customer_code', 'customer_id', 'vendor_id', 'customer_sku', 'customer_product_name', 'active', 'sort_order']),
    'prices': (m.ProductZonePrice, ['internal_product_code', 'zone', 'price', 'effective_date'],
        ['expiration_date', 'active', 'retail_price', 'return_price']),
    'price_overrides': (m.CustomerPriceOverride, ['internal_product_code', 'override_price', 'effective_date'],
        ['customer_code', 'customer_id', 'expiration_date', 'active', 'reason', 'notes']),
    'delivery_rules': (m.CustomerDeliveryRule, ['location_code', 'delivery_days'],
        ['customer_code', 'customer_id', 'production_lead_days', 'order_cutoff_time', 'timezone', 'late_order_policy']),
}
CONVERTERS = {
    'active': boolean, 'pricing_zone': integer, 'pack_size': integer, 'sort_order': integer,
    'production_lead_days': integer, 'zone': integer,
    'price': money, 'override_price': money, 'effective_date': calendar_date,
    'expiration_date': calendar_date,
    'order_cutoff_time': clock_time, 'delivery_days': weekdays,
    'units_per_tray': integer, 'retail_price': money, 'return_price': money, 'customer_id': integer,
}
TEXT_IDENTIFIERS = {'customer_code', 'location_code', 'internal_product_code', 'vendor_id', 'customer_sku', 'sku', 'upc', 'store_number', 'zip', 'phone'}


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as source:
            rows = list(csv.reader(source))
    elif path.suffix.lower() == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = []
            for cells in workbook.active.iter_rows():
                for cell in cells:
                    if cell.data_type == 'f':
                        raise ImportFailure(f'{cell.coordinate}: formulas are not accepted; supply explicit values.')
                rows.append([cell.value for cell in cells])
        finally:
            workbook.close()
    else:
        raise ImportFailure('Supported files are .csv and .xlsx.')
    if not rows:
        raise ImportFailure('File is empty.')
    header = [str(value or '').strip() for value in rows[0]]
    if not all(header) or len(header) != len(set(header)):
        raise ImportFailure('Headers must be nonempty and unique.')
    output = []
    for number, row in enumerate(rows[1:], 2):
        if all(value in (None, '') for value in row):
            continue
        if len(row) > len(header) and any(value not in (None, '') for value in row[len(header):]):
            raise ImportFailure(f'Row {number}: more values than headers.')
        padded = list(row) + [None] * max(0, len(header) - len(row))
        output.append((number, dict(zip(header, padded))))
    return header, output


def referenced(model, **lookup):
    try:
        return model.objects.get(**lookup)
    except model.DoesNotExist:
        raise ValueError(f'Unknown {model._meta.verbose_name}: {lookup}. Import its master record first.') from None


def parse_row(row, required):
    result = {}
    for field, value in row.items():
        if value in (None, ''):
            if field in required:
                raise ValueError(f'{field} is required.')
            continue
        if field in TEXT_IDENTIFIERS and not isinstance(value, str):
            raise ValueError(f'{field} must be stored as text in Excel to preserve leading zeros.')
        result[field] = CONVERTERS.get(field, lambda val: str(val).strip())(value)
    return result


def identify(kind, values):
    if kind == 'customers':
        customer_id = values.pop('customer_id', None)
        if customer_id is not None:
            customer = referenced(m.Customer, pk=customer_id)
            return {'pk': customer.pk}
        code = values.pop('customer_code', None)
        if not code:
            raise ValueError('Supply customer_id to update an existing customer, or customer_code for a keyed import. Add new customers without codes in the Customers screen.')
        return {'customer_code': code}
    if kind == 'products':
        return {'internal_product_code': values.pop('internal_product_code')}
    customer = None
    product = None
    if kind in ('locations', 'customer_products', 'price_overrides', 'delivery_rules'):
        code = values.pop('customer_code', None)
        customer_id = values.pop('customer_id', None)
        if customer_id is not None:
            customer = referenced(m.Customer, pk=customer_id)
            if code and customer.customer_code != code:
                raise ValueError('customer_id and customer_code refer to different customers.')
        elif code:
            customer = referenced(m.Customer, customer_code=code)
        else:
            raise ValueError('Supply customer_id or customer_code to identify the customer.')
    if 'internal_product_code' in values:
        product = referenced(m.Product, internal_product_code=values.pop('internal_product_code'))
    if kind == 'locations':
        return {'customer': customer, 'location_code': values.pop('location_code')}
    if kind == 'customer_products':
        return {'customer': customer, 'product': product}
    if kind == 'prices':
        return {'product': product, 'zone': values.pop('zone'), 'effective_date': values.pop('effective_date')}
    if kind == 'price_overrides':
        if not product_is_available(customer, product):
            raise ValueError('Choose an active product available in this customer\'s catalog. If the customer has a custom product list, enable the product there first.')
        return {'customer': customer, 'product': product, 'effective_date': values.pop('effective_date')}
    code = values.pop('location_code')
    location = referenced(m.CustomerLocation, customer=customer, location_code=code)
    return {'customer': customer, 'location': location}


def import_master_data(kind, path, dry_run=False, source_name=None):
    model, required, optional = SCHEMAS[kind]
    header, rows = read_rows(path)
    missing = set(required) - set(header)
    unknown = set(header) - set(required + optional)
    if missing or unknown:
        raise ImportFailure(f'Missing columns: {sorted(missing)}; unrecognized columns: {sorted(unknown)}.')
    report = ImportReport()
    errors = []
    seen = set()
    with transaction.atomic():
        for row_number, row in rows:
            try:
                with transaction.atomic():
                    values = parse_row(row, required)
                    lookup = identify(kind, values)
                    identity = tuple((key, str(value.pk) if isinstance(value, m.models.Model) else str(value)) for key, value in lookup.items())
                    if identity in seen:
                        raise ValueError('Duplicate record key in this file.')
                    seen.add(identity)
                    obj = model.objects.filter(**lookup).first()
                    before = snapshot(obj) if obj else {}
                    obj = obj or model(**lookup)
                    for field, value in values.items():
                        setattr(obj, field, value)
                    obj.full_clean()
                    if before and before == snapshot(obj):
                        report.skipped += 1
                        continue
                    is_new = obj.pk is None
                    obj.save()
                    record_change(obj, before, source=f'import:{kind}:{Path(source_name or path).name}'[:255])
                    if is_new:
                        report.inserted += 1
                    else:
                        report.updated += 1
            except (ValidationError, ValueError, InvalidOperation) as exc:
                errors.append(f'Row {row_number}: {exc}')
        if errors:
            raise ImportFailure('No records committed.\n' + '\n'.join(errors))
        if dry_run:
            transaction.set_rollback(True)
    return report
