"""Import the explicit Nelson legacy workbook layout without altering its source."""
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from bakery.models import Product, ProductSourceRecord, ProductZonePrice
from bakery.services.audit import record_change, snapshot
from bakery.services.imports import ImportFailure, ImportReport, integer, money

HEADERS = ['VendorsId', 'Vendor', 'Name', 'Sku', 'Size', 'Category', 'Price', 'Price2', 'Price3',
           'RetailPrice', 'Retail2', 'Retail3', 'ReturnPrice', 'ReturnComm', 'ReturnPrice2',
           'ReturnPrice3', 'ReturnSamePrice', 'CostPrice', 'UnitsPerTray', 'Updated', 'UserID']
ZONE_COLUMNS = {1: ('G', 'J', 'M'), 2: ('H', 'K', 'O'), 3: ('I', 'L', 'P')}


def source_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def read_nelson(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'dbo_Products' not in workbook.sheetnames:
            raise ImportFailure('Expected the dbo_Products worksheet.')
        sheet = workbook['dbo_Products']
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=21))]
        if header != HEADERS:
            raise ImportFailure('Nelson workbook headers differ from the verified source layout.')
        records = []
        keys = set()
        for row in sheet.iter_rows(min_row=2, max_col=21):
            if all(cell.value in (None, '') for cell in row):
                continue
            number = row[0].row
            if any(cell.data_type == 'f' for cell in row):
                raise ImportFailure(f'Row {number}: formulas must be replaced by explicit values.')
            values = {name: source_value(cell.value) for name, cell in zip(HEADERS, row)}
            for field in ('VendorsId', 'Sku'):
                if not isinstance(values[field], str) or not values[field].strip():
                    raise ImportFailure(f'Row {number}: {field} must be nonempty text to preserve identifiers.')
            if not values['Name']:
                raise ImportFailure(f'Row {number}: product name is required.')
            key = f"{values['Vendor'] or ''}|{values['VendorsId']}|{values['Sku']}"
            if key in keys:
                raise ImportFailure(f'Row {number}: duplicate source identity {key}.')
            keys.add(key)
            try:
                tray = integer(values['UnitsPerTray'])
                if tray < 1:
                    raise ValueError('UnitsPerTray must be positive.')
                prices = {zone: tuple(money(row[ord(column) - ord('A')].value) for column in columns)
                          for zone, columns in ZONE_COLUMNS.items()}
                if any(price < 0 for group in prices.values() for price in group):
                    raise ValueError('Prices cannot be negative.')
            except (ValueError, ArithmeticError) as exc:
                raise ImportFailure(f'Row {number}: {exc}') from exc
            records.append({'row': number, 'key': key, 'values': values, 'tray': tray, 'prices': prices})
        if not records:
            raise ImportFailure('No product rows found.')
        return records
    finally:
        workbook.close()


def import_nelson(path, effective_date, dry_run=False, source_name=None, actor=None):
    records = read_nelson(path)
    duplicate_skus = {sku for sku, count in Counter(row['values']['Sku'] for row in records).items() if count > 1}
    report = ImportReport()
    source_name = Path(source_name or path).name
    with transaction.atomic():
        for row in records:
            values = row['values']
            source = ProductSourceRecord.objects.select_related('product').filter(source_key=row['key']).first()
            # A source-derived compound key preserves duplicate SKU rows without assigning new source SKUs.
            code = f"{values['Sku']}-{values['VendorsId']}"
            product = source.product if source else Product(internal_product_code=code)
            before = snapshot(product) if product.pk else {}
            for field, value in {
                'sku': values['Sku'], 'upc': values['Sku'], 'product_name': values['Name'],
                'size': str(values['Size'] or ''), 'category': values['Category'] or '',
                'units_per_tray': row['tray'], 'sort_order': row['row'] - 2,
            }.items():
                setattr(product, field, value)
            if not source:
                product.pack_size = 1
                product.unit_of_measure = 'item'
            product.full_clean()
            changed = not before or before != snapshot(product)
            if changed:
                product.save()
                record_change(product, before, f'nelson:{source_name}', actor)
            notes = 'Duplicate source SKU: verify barcode before scanning or merging.' if values['Sku'] in duplicate_skus else ''
            source = source or ProductSourceRecord(product=product, source_key=row['key'])
            source_before = snapshot(source) if source.pk else {}
            source.workbook, source.sheet, source.row_number = source_name, 'dbo_Products', row['row']
            source.values, source.review_notes = values, notes
            source.full_clean()
            if not source.pk or source_before != snapshot(source):
                source.save()
                record_change(source, source_before, f'nelson:{source_name}', actor)
                changed = True
            for zone, (wholesale, retail, returned) in row['prices'].items():
                price = ProductZonePrice.objects.filter(product=product, zone=zone, effective_date=effective_date).first()
                price_before = snapshot(price) if price else {}
                price = price or ProductZonePrice(product=product, zone=zone, effective_date=effective_date)
                price.price, price.retail_price, price.return_price = wholesale, retail, returned
                price.full_clean()
                if not price.pk or price_before != snapshot(price):
                    price.save()
                    record_change(price, price_before, f'nelson:{source_name}', actor)
                    changed = True
            if not before:
                report.inserted += 1
            elif changed:
                report.updated += 1
            else:
                report.skipped += 1
        if dry_run:
            transaction.set_rollback(True)
    return report, sorted(duplicate_skus)
