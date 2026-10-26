import csv
import tempfile
from datetime import date, datetime, time
from decimal import Decimal
from io import StringIO
from pathlib import Path
from zoneinfo import ZoneInfo

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import Workbook

from bakery import models as m
from bakery.services.access import customers_for, locations_for, orders_for, products_for
from bakery.services.delivery import available_delivery_dates, required_production_date, rules_for, validate_delivery_date
from bakery.services.imports import ImportFailure, import_master_data
from bakery.services.pricing import PricingBook, resolve_price


class FoundationFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.customer = m.Customer.objects.create(customer_code='A001', customer_name='Test account A', pricing_zone=2)
        cls.other = m.Customer.objects.create(customer_code='B001', customer_name='Private account B', pricing_zone=1)
        cls.location = m.CustomerLocation.objects.create(customer=cls.customer, location_code='01', location_name='Test store A')
        cls.second_location = m.CustomerLocation.objects.create(customer=cls.customer, location_code='02', location_name='Private sibling store')
        cls.other_location = m.CustomerLocation.objects.create(customer=cls.other, location_code='01', location_name='Private store B')
        cls.product = m.Product.objects.create(internal_product_code='00017', product_name='Test bread', upc='000123456789', pack_size=6)
        cls.restricted = m.Product.objects.create(internal_product_code='00018', product_name='Private pastry')
        cls.link = m.CustomerProduct.objects.create(customer=cls.customer, product=cls.product, customer_sku='0009', vendor_id='0081')
        m.CustomerProduct.objects.create(customer=cls.other, product=cls.restricted)
        cls.price = m.ProductZonePrice.objects.create(product=cls.product, zone=2, price='21.00', effective_date=date(2026, 1, 1))
        cls.rule = m.CustomerDeliveryRule.objects.create(customer=cls.customer, location=cls.location, delivery_days=[1, 3, 5])
        cls.user = get_user_model().objects.create_user('test-customer', password='Test-only-password-493!')
        m.CustomerAccess.objects.create(user=cls.user, customer=cls.customer, location=cls.location)
        cls.now = datetime(2026, 10, 5, 13, 59, tzinfo=ZoneInfo('America/New_York'))


class SchemaTests(FoundationFixture):
    def test_customer_zone_constraint(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            m.Customer.objects.create(customer_code='bad', customer_name='Invalid', pricing_zone=4)

    def test_location_codes_are_unique_per_customer(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            m.CustomerLocation.objects.create(customer=self.customer, location_code='01', location_name='Duplicate')

    def test_access_cannot_cross_customer(self):
        with self.assertRaises(ValidationError):
            m.CustomerAccess(user=self.user, customer=self.customer, location=self.other_location).full_clean()

    def test_delivery_rule_cannot_cross_customer(self):
        with self.assertRaises(ValidationError):
            m.CustomerDeliveryRule(customer=self.customer, location=self.other_location, delivery_days=[1]).full_clean()

    def test_invalid_delivery_days(self):
        for days in ([], [7], ['1'], [1, 1], [True]):
            with self.subTest(days=days), self.assertRaises(ValidationError):
                m.validate_weekdays(days)

    def test_bad_timezone(self):
        with self.assertRaises(ValidationError):
            m.validate_timezone('not/a-timezone')

    def test_order_customer_location_must_match(self):
        order = m.Order(customer=self.customer, customer_location=self.other_location)
        with self.assertRaises(ValidationError):
            order.clean()

    def test_price_period_validation(self):
        price = m.ProductZonePrice(product=self.product, zone=1, price=1,
            effective_date=date(2026, 2, 1), expiration_date=date(2026, 1, 1))
        with self.assertRaises(ValidationError):
            price.full_clean()


class PricingTests(FoundationFixture):
    def test_assigned_zone(self):
        result = resolve_price(self.customer, self.product, date(2026, 10, 3))
        self.assertEqual(result.price, Decimal('21.00'))
        self.assertEqual(result.price_source, 'ZONE_2')

    def test_override_wins(self):
        m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
            override_price='19.75', effective_date=date(2026, 10, 1))
        result = resolve_price(self.customer, self.product, date(2026, 10, 3))
        self.assertEqual((result.price, result.price_source, result.standard_price), (Decimal('19.75'), 'CUSTOMER_OVERRIDE', Decimal('21.00')))

    def test_zero_override_is_valid(self):
        m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
            override_price=0, effective_date=date(2026, 10, 1))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('0'))

    def test_expired_future_and_inactive_overrides_do_not_apply(self):
        for start, end, active in [(date(2026, 1, 1), date(2026, 9, 1), True),
                (date(2026, 11, 1), None, True), (date(2026, 10, 1), None, False)]:
            m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
                override_price=7, effective_date=start, expiration_date=end, active=active)
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price_source, 'ZONE_2')

    def test_inclusive_price_period_end(self):
        m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
            override_price=5, effective_date=date(2026, 10, 1), expiration_date=date(2026, 10, 3))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('5'))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 4)).price, Decimal('21'))

    def test_missing_assigned_zone_does_not_fall_back(self):
        self.price.delete()
        m.ProductZonePrice.objects.create(product=self.product, zone=1, price=20, effective_date=date(2026, 1, 1))
        result = resolve_price(self.customer, self.product, date(2026, 10, 3))
        self.assertIsNone(result.price)
        self.assertEqual(result.price_source, 'MISSING')

    def test_product_availability_enforced(self):
        with self.assertRaises(PermissionDenied):
            resolve_price(self.customer, self.restricted, date(2026, 10, 3))
        self.link.active = False
        self.link.save()
        with self.assertRaises(PermissionDenied):
            resolve_price(self.customer, self.product, date(2026, 10, 3))

    def test_inactive_customer(self):
        self.customer.active = False
        with self.assertRaises(PermissionDenied):
            resolve_price(self.customer, self.product, date(2026, 10, 3))

    def test_overlapping_period_rejected(self):
        price = m.ProductZonePrice(product=self.product, zone=2, price=22, effective_date=date(2026, 10, 1))
        with self.assertRaises(ValidationError):
            price.full_clean()

    def test_corrupt_overlap_fails_closed(self):
        m.ProductZonePrice.objects.create(product=self.product, zone=2, price=22, effective_date=date(2026, 10, 1))
        with self.assertRaises(ValidationError):
            resolve_price(self.customer, self.product, date(2026, 10, 3))

    def test_catalog_pricing_has_constant_query_count(self):
        with self.assertNumQueries(3):
            book = PricingBook(self.customer, date(2026, 10, 3))
            for _ in range(20):
                book.resolve(self.product)

    def test_price_change_preserves_historical_line(self):
        order = m.Order.objects.create(order_number='TEST-ONLY-1', customer=self.customer,
            customer_location=self.location, order_date=date(2026, 10, 3), requested_delivery_date=date(2026, 10, 6))
        item = m.OrderItem.objects.create(order=order, product=self.product, quantity=2, unit_price=21,
            line_total=42, price_source='ZONE_2', product_name='Test bread', unit_of_measure='case', required_production_date=date(2026, 10, 5))
        self.price.price = 22
        self.price.save()
        item.refresh_from_db()
        self.assertEqual(item.unit_price, Decimal('21'))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('22'))


class DeliveryTests(FoundationFixture):
    def test_before_cutoff(self):
        result = validate_delivery_date(self.customer, self.location, date(2026, 10, 6), self.now)
        self.assertFalse(result.late)
        self.assertEqual(result.cutoff.hour, 14)

    def test_exact_cutoff_requires_approval(self):
        result = validate_delivery_date(self.customer, self.location, date(2026, 10, 6), self.now.replace(hour=14, minute=0))
        self.assertTrue(result.late)
        self.assertTrue(result.requires_approval)

    def test_late_allow_and_prohibit(self):
        self.rule.late_order_policy = m.LatePolicy.ALLOW
        self.rule.save()
        result = validate_delivery_date(self.customer, self.location, date(2026, 10, 6), self.now.replace(hour=15))
        self.assertTrue(result.late)
        self.assertFalse(result.requires_approval)
        self.rule.late_order_policy = m.LatePolicy.PROHIBIT
        self.rule.save()
        with self.assertRaises(ValidationError):
            validate_delivery_date(self.customer, self.location, date(2026, 10, 6), self.now.replace(hour=15))

    def test_non_delivery_and_past_dates_rejected(self):
        for day in (date(2026, 10, 7), date(2026, 10, 3), date(2026, 10, 5)):
            with self.subTest(day=day), self.assertRaises(ValidationError):
                validate_delivery_date(self.customer, self.location, day, self.now)

    def test_store_rules_do_not_inherit_customer_or_sibling_settings(self):
        self.rule.delivery_days = [4]
        self.rule.production_lead_days = 3
        self.rule.order_cutoff_time = time(12)
        self.rule.save()
        self.customer.timezone = 'America/Los_Angeles'
        with self.assertRaises(ValidationError):
            rules_for(self.customer, self.second_location)
        result = validate_delivery_date(self.customer, self.location, date(2026, 10, 9), self.now)
        self.assertEqual(result.cutoff, datetime(2026, 10, 6, 12, tzinfo=ZoneInfo('America/New_York')))

    def test_missing_calendar_is_not_invented(self):
        self.rule.delete()
        with self.assertRaises(ValidationError):
            available_delivery_dates(self.customer, self.location, self.now)

    def test_available_dates_skip_prohibited_deadlines(self):
        self.rule.late_order_policy = m.LatePolicy.PROHIBIT
        self.rule.save()
        dates = available_delivery_dates(self.customer, self.location, self.now.replace(hour=15), horizon=7)
        self.assertEqual([entry.delivery_date for entry in dates], [date(2026, 10, 8), date(2026, 10, 10)])

    def test_timezone_and_dst(self):
        self.rule.timezone = 'America/Los_Angeles'
        self.rule.save()
        result = validate_delivery_date(self.customer, self.location, date(2026, 11, 3),
            datetime(2026, 11, 2, 20, tzinfo=ZoneInfo('UTC')))
        self.assertFalse(result.late)
        self.assertEqual(result.cutoff.utcoffset().total_seconds(), -8 * 3600)

    def test_naive_time_is_rejected(self):
        with self.assertRaises(ValidationError):
            validate_delivery_date(self.customer, self.location, date(2026, 10, 6), datetime(2026, 10, 5, 13))

    def test_product_production_date_is_separate(self):
        self.product.production_lead_days = 2
        self.assertEqual(required_production_date(self.product, date(2026, 10, 6)), date(2026, 10, 4))


class ImportTests(FoundationFixture):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)

    def file(self, text):
        path = Path(self.temp.name) / 'input.csv'
        path.write_text(text, encoding='utf-8')
        return path

    def test_customer_import_insert_update_skip_and_audit(self):
        path = self.file('customer_code,customer_name,pricing_zone\n0007,Imported account,3\n')
        self.assertEqual(import_master_data('customers', path).inserted, 1)
        self.assertEqual(import_master_data('customers', path).skipped, 1)
        path = self.file('customer_code,customer_name,pricing_zone\n0007,Updated account,2\n')
        self.assertEqual(import_master_data('customers', path).updated, 1)
        self.assertEqual(m.MasterDataChange.objects.count(), 2)
        self.assertEqual(m.Customer.objects.get(customer_code='0007').pricing_zone, 2)

    def test_error_rolls_back_whole_file(self):
        path = self.file('customer_code,customer_name,pricing_zone\n0007,Valid,1\n0008,Invalid,4\n')
        with self.assertRaisesRegex(ImportFailure, 'Row 3'):
            import_master_data('customers', path)
        self.assertFalse(m.Customer.objects.filter(customer_code='0007').exists())
        self.assertEqual(m.MasterDataChange.objects.count(), 0)

    def test_dry_run_does_not_save_or_audit(self):
        path = self.file('customer_code,customer_name,pricing_zone\n0007,Valid,1\n')
        self.assertEqual(import_master_data('customers', path, dry_run=True).inserted, 1)
        self.assertFalse(m.Customer.objects.filter(customer_code='0007').exists())
        self.assertFalse(m.MasterDataChange.objects.exists())

    def test_unknown_customer_mapping_is_reported(self):
        path = self.file('customer_code,internal_product_code\nUnknown,00017\n')
        with self.assertRaisesRegex(ImportFailure, 'Unknown customer'):
            import_master_data('customer_products', path)

    def test_duplicate_keys_are_rejected(self):
        path = self.file('customer_code,customer_name,pricing_zone\n0007,First,1\n0007,Second,2\n')
        with self.assertRaisesRegex(ImportFailure, 'Duplicate record key'):
            import_master_data('customers', path)

    def test_unknown_columns_are_not_silently_lost(self):
        path = self.file('customer_code,customer_name,pricing_zone,mystery\n0007,First,1,Important\n')
        with self.assertRaisesRegex(ImportFailure, 'mystery'):
            import_master_data('customers', path)

    def test_duplicate_headers_are_rejected(self):
        path = self.file('customer_code,customer_code,pricing_zone\n0007,First,1\n')
        with self.assertRaisesRegex(ImportFailure, 'unique'):
            import_master_data('customers', path)

    def test_upc_and_sku_leading_zeros_survive(self):
        path = self.file('internal_product_code,product_name,pack_size,unit_of_measure,upc\n00019,Imported bread,6,case,000123450000\n')
        import_master_data('products', path)
        product = m.Product.objects.get(internal_product_code='00019')
        self.assertEqual(product.upc, '000123450000')

    def test_numeric_excel_identifier_rejected(self):
        workbook = Workbook()
        workbook.active.append(['internal_product_code', 'product_name', 'pack_size', 'unit_of_measure', 'upc'])
        workbook.active.append(['00019', 'Imported bread', 6, 'case', 12345])
        path = Path(self.temp.name) / 'input.xlsx'
        workbook.save(path)
        with self.assertRaisesRegex(ImportFailure, 'upc must be stored as text'):
            import_master_data('products', path)

    def test_excel_text_and_native_date_import(self):
        workbook = Workbook()
        workbook.active.append(['internal_product_code', 'zone', 'price', 'effective_date'])
        workbook.active.append(['00017', 1, 19.5, datetime(2026, 1, 1)])
        path = Path(self.temp.name) / 'input.xlsx'
        workbook.save(path)
        self.assertEqual(import_master_data('prices', path).inserted, 1)

    def test_excel_formula_rejected(self):
        workbook = Workbook()
        workbook.active.append(['customer_code', 'customer_name', 'pricing_zone'])
        workbook.active.append(['0009', '=1+1', 1])
        path = Path(self.temp.name) / 'input.xlsx'
        workbook.save(path)
        with self.assertRaisesRegex(ImportFailure, 'formulas'):
            import_master_data('customers', path)

    def test_fractional_pack_and_price_precision_rejected(self):
        path = self.file('internal_product_code,product_name,pack_size,unit_of_measure\n00019,Bread,1.5,case\n')
        with self.assertRaises(ImportFailure):
            import_master_data('products', path)
        path = self.file('internal_product_code,zone,price,effective_date\n00017,1,1.999,2026-01-01\n')
        with self.assertRaises(ImportFailure):
            import_master_data('prices', path)

    def test_locations_relationships_overrides_and_delivery_import(self):
        import_master_data('locations', self.file('customer_code,location_code,location_name\nA001,003,Imported store\n'))
        import_master_data('customer_products', self.file('customer_code,internal_product_code,vendor_id\nA001,00018,000033\n'))
        import_master_data('price_overrides', self.file('customer_code,internal_product_code,override_price,effective_date\nA001,00018,12.00,2026-01-01\n'))
        import_master_data('delivery_rules', self.file('customer_code,location_code,delivery_days,production_lead_days\nA001,003,"1,3",2\n'))
        self.assertEqual(m.CustomerProduct.objects.get(customer=self.customer, product=self.restricted).vendor_id, '000033')
        self.assertEqual(m.CustomerDeliveryRule.objects.get(location__location_code='003').delivery_days, [1, 3])

    def test_full_catalog_price_override_import_needs_no_product_mapping(self):
        self.link.delete()
        path = self.file('customer_code,internal_product_code,override_price,effective_date\nA001,00017,12.00,2026-01-01\n')
        report = import_master_data('price_overrides', path)
        self.assertEqual(report.inserted, 1)
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('12'))
        self.assertFalse(m.CustomerProduct.objects.filter(customer=self.customer).exists())
        self.assertEqual(import_master_data('price_overrides', path).skipped, 1)

    def test_override_import_preserves_custom_catalog_exclusions(self):
        path = self.file('customer_code,internal_product_code,override_price,effective_date\nA001,00018,12.00,2026-01-01\n')
        with self.assertRaisesRegex(ImportFailure, 'enable the product there first'):
            import_master_data('price_overrides', path)
        self.link.active = False
        self.link.save()
        path = self.file('customer_code,internal_product_code,override_price,effective_date\nA001,00017,12.00,2026-01-01\n')
        with self.assertRaisesRegex(ImportFailure, 'enable the product there first'):
            import_master_data('price_overrides', path)
        self.assertFalse(m.CustomerPriceOverride.objects.exists())

    def test_blank_optional_values_leave_existing_values(self):
        import_master_data('customer_products', self.file('customer_code,internal_product_code,vendor_id\nA001,00017,\n'))
        self.link.refresh_from_db()
        self.assertEqual(self.link.vendor_id, '0081')

    def test_import_price_overlap_rolls_back(self):
        path = self.file('internal_product_code,zone,price,effective_date\n00017,2,23,2026-10-01\n')
        with self.assertRaisesRegex(ImportFailure, 'overlap'):
            import_master_data('prices', path)

    def test_price_import_is_idempotent_for_integer_and_decimal_values(self):
        path = self.file('internal_product_code,zone,price,effective_date\n00017,1,19,2026-01-01\n')
        self.assertEqual(import_master_data('prices', path).inserted, 1)
        self.assertEqual(import_master_data('prices', path).skipped, 1)
        path = self.file('internal_product_code,zone,price,effective_date\n00017,1,19.00,2026-01-01\n')
        self.assertEqual(import_master_data('prices', path).skipped, 1)


class AccessTests(FoundationFixture):
    def setUp(self):
        self.client.force_login(self.user)

    def test_customer_only_sees_assigned_location(self):
        self.assertEqual(list(customers_for(self.user)), [self.customer])
        self.assertEqual(list(locations_for(self.user)), [self.location])
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, 'Test store A')
        self.assertNotContains(response, 'Private sibling store')
        self.assertNotContains(response, 'Private account B')

    def test_direct_customer_catalog_access_is_scoped(self):
        self.assertEqual(self.client.get(reverse('catalog', args=[self.other.pk])).status_code, 404)
        with self.assertRaises(PermissionDenied):
            products_for(self.user, self.other)

    def test_customer_catalog_hides_restricted_products_and_prices(self):
        response = self.client.get(reverse('catalog', args=[self.customer.pk]))
        self.assertContains(response, 'Test bread')
        self.assertContains(response, '$21.00')
        self.assertNotContains(response, 'Private pastry')

    def test_search_by_customer_sku_and_vendor_id(self):
        for query in ('0009', '0081', '000123456789', '00017'):
            self.assertContains(self.client.get(reverse('catalog', args=[self.customer.pk]), {'q': query}), 'Test bread')

    def test_customer_cannot_access_office_endpoints(self):
        for name in ('customers', 'products', 'pricing', 'price-records', 'special-prices', 'imports'):
            with self.subTest(name=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 403)
        self.assertEqual(self.client.get(reverse('import-template', args=['prices'])).status_code, 403)
        self.assertEqual(self.client.post(reverse('imports'), {}).status_code, 403)

    def test_staff_flag_does_not_grant_office_access(self):
        self.user.is_staff = True
        self.user.save()
        self.assertEqual(self.client.get(reverse('customers')).status_code, 403)
        self.assertEqual(self.client.get(reverse('admin:bakery_customer_changelist')).status_code, 403)

    def test_access_requires_explicit_store_assignments(self):
        new_store = m.CustomerLocation.objects.create(customer=self.customer, location_code='NEW', location_name='New store')
        self.assertEqual(list(locations_for(self.user)), [self.location])
        m.CustomerAccess.objects.create(user=self.user, customer=self.customer, location=self.second_location)
        self.assertEqual(locations_for(self.user).count(), 2)
        self.assertFalse(locations_for(self.user).filter(pk=new_store.pk).exists())
        with self.assertRaises(ValidationError):
            m.CustomerAccess(user=self.user, customer=self.customer).full_clean()
        with self.assertRaises(IntegrityError), transaction.atomic():
            m.CustomerAccess.objects.create(user=self.user, customer=self.customer)

    def test_inactive_location_and_customer_assignments_hidden(self):
        self.location.active = False
        self.location.save()
        self.assertFalse(locations_for(self.user).exists())
        self.assertFalse(customers_for(self.user).exists())

    def test_bad_cross_customer_assignment_fails_closed(self):
        m.CustomerAccess.objects.filter(user=self.user).delete()
        m.CustomerAccess.objects.create(user=self.user, customer=self.customer, location=self.other_location)
        self.assertFalse(customers_for(self.user).exists())

    def test_orders_scope_excludes_other_locations(self):
        for number, customer, location in [('TA1', self.customer, self.location), ('TA2', self.customer, self.second_location), ('TB1', self.other, self.other_location)]:
            m.Order.objects.create(order_number=number, customer=customer, customer_location=location,
                order_date=date(2026, 10, 3), requested_delivery_date=date(2026, 10, 6))
        self.assertEqual(list(orders_for(self.user).values_list('order_number', flat=True)), ['TA1'])

    def test_anonymous_requires_login(self):
        self.client.logout()
        self.assertRedirects(self.client.get(reverse('dashboard')), '/accounts/login/?next=/')

    def test_login_and_post_only_logout(self):
        self.client.logout()
        self.assertRedirects(self.client.post(reverse('login'), {'username': 'test-customer', 'password': 'Test-only-password-493!'}), '/')
        self.assertEqual(self.client.get(reverse('logout')).status_code, 405)
        self.assertRedirects(self.client.post(reverse('logout')), reverse('login'))

    def test_csrf_required_for_login(self):
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.post(reverse('login'), {'username': 'x', 'password': 'x'}).status_code, 403)

    def test_production_role_cannot_manage_master_data(self):
        call_command('setup_roles', stdout=StringIO())
        self.user.groups.add(Group.objects.get(name='PRODUCTION'))
        self.assertRedirects(self.client.get(reverse('dashboard')), reverse('bake-plan'))
        self.assertEqual(self.client.get(reverse('pricing')).status_code, 403)

    def test_office_role_and_repeatable_role_setup(self):
        call_command('setup_roles', stdout=StringIO())
        call_command('setup_roles', stdout=StringIO())
        self.assertEqual(Group.objects.count(), 4)
        self.user.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        self.assertContains(self.client.get(reverse('customers')), 'Private account B')
        self.assertContains(self.client.get(reverse('dashboard')), 'setup checklist')

    def test_office_import_form_and_template(self):
        call_command('setup_roles', stdout=StringIO())
        self.user.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        self.assertEqual(self.client.get(reverse('imports')).status_code, 200)
        response = self.client.get(reverse('import-template', args=['customers']))
        self.assertIn('customer_name,pricing_zone,customer_code,customer_id', response.content.decode())
        upload = SimpleUploadedFile('customers.csv', b'customer_code,customer_name,pricing_zone\nUI1,Imported through UI,1\n')
        self.assertContains(self.client.post(reverse('imports'), {'kind': 'customers', 'file': upload, 'dry_run': 'on'}), 'Validated without saving')
        self.assertFalse(m.Customer.objects.filter(customer_code='UI1').exists())
        upload = SimpleUploadedFile('customers.csv', b'customer_code,customer_name,pricing_zone\nUI1,Imported through UI,1\n')
        self.assertContains(self.client.post(reverse('imports'), {'kind': 'customers', 'file': upload}), 'Imported successfully')
        self.assertTrue(m.Customer.objects.filter(customer_code='UI1').exists())

    def test_office_pages_render_and_empty_search(self):
        call_command('setup_roles', stdout=StringIO())
        self.user.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        for name in ('dashboard', 'customers', 'products', 'price-records', 'special-prices', 'imports', 'account'):
            self.assertEqual(self.client.get(reverse(name), {'q': 'unlikely-result'}).status_code, 200)

    def test_customer_can_change_own_password(self):
        self.assertEqual(self.client.get(reverse('password-change')).status_code, 200)
        response = self.client.post(reverse('password-change'), {
            'old_password': 'Test-only-password-493!',
            'new_password1': 'New-test-only-password-887!',
            'new_password2': 'New-test-only-password-887!',
        })
        self.assertRedirects(response, reverse('password-change-done'))
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password('New-test-only-password-887!'))
        self.assertEqual(self.client.get(reverse('dashboard')).status_code, 200)

    def test_bad_pricing_shows_review_message(self):
        m.ProductZonePrice.objects.create(product=self.product, zone=2, price=22, effective_date=date(2026, 2, 1))
        response = self.client.get(reverse('catalog', args=[self.customer.pk]), follow=True)
        self.assertContains(response, 'Your pricing needs bakery office review')

    def test_uploaded_source_filename_is_audited(self):
        call_command('setup_roles', stdout=StringIO())
        self.user.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        upload = SimpleUploadedFile('real-customer-list.csv', b'customer_code,customer_name,pricing_zone\nUI2,Another import,1\n')
        self.client.post(reverse('imports'), {'kind': 'customers', 'file': upload})
        self.assertEqual(m.MasterDataChange.objects.get().source, 'import:customers:real-customer-list.csv')
