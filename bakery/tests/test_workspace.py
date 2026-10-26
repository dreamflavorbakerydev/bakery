import json
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.test import TestCase, Client
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from bakery import models as m
from bakery.services.imports import ImportFailure
from bakery.services.nelson import HEADERS, import_nelson
from bakery.services.pricing import resolve_price
from bakery.tests.test_foundation import FoundationFixture


class NelsonTests(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'DFB_Products_Nelson.xlsx'

    def workbook(self, rows=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'dbo_Products'
        sheet.append(HEADERS)
        rows = rows or [
            ['00001', 'Seabra', 'Sample cake', '000123456789', '14 oz', 'Cake',
             '10', '11', '12', '15', '16', '17', '8', '7', '9', '10', True, '5', '12', '2026-01-01', 'Source user']]
        for row in rows:
            sheet.append(row)
        workbook.save(self.path)
        workbook.close()

    def test_all_nine_price_columns_and_identifiers_preserved(self):
        self.workbook()
        report, duplicates = import_nelson(self.path, date(2026, 10, 3))
        self.assertEqual(report.inserted, 1)
        self.assertEqual(duplicates, [])
        product = m.Product.objects.get()
        self.assertEqual(product.sku, '000123456789')
        self.assertEqual(product.upc, '000123456789')
        self.assertEqual(product.size, '14 oz')
        self.assertEqual(product.units_per_tray, 12)
        self.assertEqual((product.pack_size, product.unit_of_measure), (1, 'item'))
        self.assertEqual(product.source_record.values['VendorsId'], '00001')
        self.assertEqual(product.source_record.values['ReturnComm'], '7')
        self.assertEqual(list(product.zone_prices.order_by('zone').values_list('price', 'retail_price', 'return_price')),
                         [(10, 15, 8), (11, 16, 9), (12, 17, 10)])
        self.assertFalse(m.Customer.objects.exists())
        self.assertFalse(m.CustomerProduct.objects.exists())

    def test_repeated_import_is_idempotent(self):
        self.workbook()
        import_nelson(self.path, date(2026, 10, 3))
        before = m.MasterDataChange.objects.count()
        report, _ = import_nelson(self.path, date(2026, 10, 3))
        self.assertEqual((report.inserted, report.updated, report.skipped), (0, 0, 1))
        self.assertEqual(m.MasterDataChange.objects.count(), before)
        self.assertEqual(m.ProductZonePrice.objects.count(), 3)

    def test_dry_run_rolls_back(self):
        self.workbook()
        report, _ = import_nelson(self.path, date(2026, 10, 3), dry_run=True)
        self.assertEqual(report.inserted, 1)
        for model in (m.Product, m.ProductSourceRecord, m.ProductZonePrice, m.MasterDataChange):
            self.assertFalse(model.objects.exists())

    def test_shared_sku_keeps_both_products(self):
        first = ['163', '', 'Copitas', '655746112342', '5', 'Cake', '4', '4', '4', '7', '7', '7', '4', '4', '4', '4', True, '2', '6', '2026-01-01', 'Source']
        second = first.copy()
        second[0], second[2], second[6] = '164', 'Flan', '9'
        self.workbook([first, second])
        _, duplicates = import_nelson(self.path, date(2026, 10, 3))
        self.assertEqual(duplicates, ['655746112342'])
        self.assertEqual(m.Product.objects.filter(sku='655746112342').count(), 2)
        self.assertEqual(m.ProductSourceRecord.objects.exclude(review_notes='').count(), 2)

    def test_numeric_sku_is_rejected(self):
        self.workbook()
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        workbook.active['D2'] = 123456789
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ImportFailure, 'Sku must be nonempty text'):
            import_nelson(self.path, date(2026, 10, 3))
        self.assertFalse(m.Product.objects.exists())


class WorkspaceTests(FoundationFixture):
    def setUp(self):
        call_command('setup_roles', stdout=StringIO())
        self.admin_user = get_user_model().objects.create_superuser('workspace-admin', password='Test-admin-password-339!')
        self.client.force_login(self.admin_user)

    def price_post(self, day):
        data = {'effective_date': day.isoformat()}
        for zone in range(1, 4):
            data.update({f'zone_{zone}_price': str(20 + zone), f'zone_{zone}_retail_price': str(30 + zone),
                         f'zone_{zone}_return_price': str(10 + zone)})
        return data

    def test_custom_administration_and_old_admin_entry(self):
        self.assertContains(self.client.get(reverse('administration')), 'A home for the details.')
        self.assertRedirects(self.client.get('/admin/'), reverse('administration'))
        for name in ('administration', 'products', 'customers', 'price-records', 'special-prices'):
            response = self.client.get(reverse(name))
            self.assertContains(response, 'Products &amp; Pricing')
            self.assertNotContains(response, 'href="/pricing/"')
            self.assertNotContains(response, 'maintenance-admin')
            self.assertNotContains(response, '/admin/bakery/')

    def test_old_pricing_links_keep_search_and_pagination(self):
        response = self.client.get(reverse('pricing'), {'q': 'Test bread', 'page': 2})
        self.assertRedirects(response, reverse('price-records') + '?q=Test+bread&page=2')

    def test_custom_editors_and_lists_render(self):
        for kind in ('customers', 'locations', 'products', 'availability', 'overrides', 'delivery', 'access'):
            self.assertEqual(self.client.get(reverse('workspace-records', args=[kind]), follow=True).status_code, 200)
            self.assertEqual(self.client.get(reverse('workspace-record-new', args=[kind]), follow=True).status_code, 200)
        for name in ('workspace-users', 'workspace-user-create', 'workspace-audit'):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        self.assertContains(self.client.get(reverse('product-detail', args=[self.product.pk])), 'Three zones. Every price.')
        self.assertContains(self.client.get(reverse('customer-detail', args=[self.customer.pk])), 'ZONE 2')

    def test_zone_price_sort_uses_current_numeric_prices_across_pages(self):
        today = timezone.localdate()
        products = m.Product.objects.bulk_create([
            m.Product(internal_product_code=f'SORT-{number}', product_name=f'Sort probe {number:02}', category='Sort tests')
            for number in range(1, 33)])
        m.ProductZonePrice.objects.bulk_create([
            m.ProductZonePrice(product=product, zone=1, price=number, effective_date=today)
            for number, product in enumerate(products[:31], 1)])
        # Neither future nor inactive records should affect the price shown or its order.
        m.ProductZonePrice.objects.create(product=products[0], zone=1, price=9999,
                                         effective_date=today + timedelta(days=10))
        m.ProductZonePrice.objects.create(product=products[-1], zone=1, price=0, active=False,
                                         effective_date=today)
        params = {'q': 'Sort probe', 'category': 'Sort tests', 'sort': 'zone1', 'direction': 'asc'}
        response = self.client.get(reverse('products'), params)
        self.assertEqual([p.pk for p in response.context['page']], [p.pk for p in products[:30]])
        self.assertContains(response, 'sort=zone1&amp;direction=asc&amp;page=2')
        self.assertContains(response, 'Zone 1: sort descending')
        response = self.client.get(reverse('products'), {**params, 'page': 2})
        self.assertEqual([p.pk for p in response.context['page']], [p.pk for p in products[30:]])
        response = self.client.get(reverse('products'), {**params, 'direction': 'desc'})
        self.assertEqual(response.context['page'][0].pk, products[30].pk)
        self.assertEqual(response.context['page'][0].display_prices[0], Decimal('31'))
        response = self.client.get(reverse('products'), {**params, 'sort': 'unknown-field'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['sort'], 'product')

    def test_customer_zone_is_required_in_custom_form(self):
        response = self.client.post(reverse('workspace-record-new', args=['customers']), {
            'customer_code': 'MISSING', 'customer_name': 'No zone', 'default_lead_days': 1,
            'default_cutoff_time': '14:00', 'timezone': 'America/New_York', 'late_order_policy': 'APPROVAL'})
        self.assertContains(response, 'This field is required')
        self.assertFalse(m.Customer.objects.filter(customer_code='MISSING').exists())

    def test_custom_product_create_and_audit(self):
        response = self.client.post(reverse('workspace-record-new', args=['products']), {
            'internal_product_code': 'CUSTOM-001', 'sku': '000001', 'product_name': 'Custom product',
            'category': 'Bread', 'pack_size': 1, 'unit_of_measure': 'item', 'production_lead_days': 1,
            'sort_order': 0, 'units_per_tray': 12, 'active': 'on'})
        product = m.Product.objects.get(internal_product_code='CUSTOM-001')
        self.assertRedirects(response, reverse('product-detail', args=[product.pk]))
        change = m.MasterDataChange.objects.get(source='workspace')
        self.assertEqual(change.actor, self.admin_user)
        self.assertEqual(change.after['sku'], '000001')

    def test_zone_grid_saves_all_values_and_closes_prior_period(self):
        day = date(2026, 10, 3)
        self.client.post(reverse('product-prices', args=[self.product.pk]), self.price_post(day))
        self.price.refresh_from_db()
        self.assertEqual(self.price.expiration_date, day - timedelta(days=1))
        self.assertEqual(list(m.ProductZonePrice.objects.filter(product=self.product, effective_date=day).order_by('zone').values_list(
            'price', 'retail_price', 'return_price')), [(21, 31, 11), (22, 32, 12), (23, 33, 13)])
        self.assertEqual(resolve_price(self.customer, self.product, day).price, 22)

    def test_invalid_zone_grid_does_not_partially_save(self):
        data = self.price_post(date(2026, 10, 3))
        data['zone_3_price'] = '-1'
        response = self.client.post(reverse('product-prices', args=[self.product.pk]), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(m.ProductZonePrice.objects.count(), 1)
        self.price.refresh_from_db()
        self.assertIsNone(self.price.expiration_date)

    def test_future_overlap_rolls_back_entire_grid(self):
        m.ProductZonePrice.objects.create(product=self.product, zone=3, effective_date=date(2026, 11, 1), price=50)
        response = self.client.post(reverse('product-prices', args=[self.product.pk]), self.price_post(date(2026, 10, 3)))
        self.assertContains(response, 'overlap')
        self.assertFalse(m.ProductZonePrice.objects.filter(effective_date=date(2026, 10, 3)).exists())
        self.price.refresh_from_db()
        self.assertIsNone(self.price.expiration_date)

    def test_customer_special_price_wins_over_zone(self):
        response = self.client.post(reverse('workspace-record-new', args=['overrides']), {
            'customer': self.customer.pk, 'product': self.product.pk, 'override_price': '9.75',
            'effective_date': '2026-10-03', 'active': 'on', 'reason': 'Contract price'})
        self.assertRedirects(response, reverse('special-prices'))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('9.75'))

    def test_override_requires_available_customer_product(self):
        response = self.client.post(reverse('workspace-record-new', args=['overrides']), {
            'customer': self.customer.pk, 'product': self.restricted.pk, 'override_price': 5,
            'effective_date': '2026-10-03', 'active': 'on'})
        self.assertContains(response, 'enable the product there first')
        self.assertFalse(m.CustomerPriceOverride.objects.exists())

    def test_delete_special_price_requires_confirmation_and_preserves_order_prices(self):
        from bakery.services.orders import create_draft, submit_order
        price = m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
            override_price='9.75', effective_date=date(2026, 1, 1))
        with patch('django.utils.timezone.now', return_value=self.now):
            order = create_draft(self.user, self.location.pk, date(2026, 10, 6), {self.product.pk: 2})
            submit_order(self.user, order.pk)
        url = reverse('special-price-delete', args=[price.pk])
        self.assertContains(self.client.get(reverse('special-prices')), url)
        self.assertContains(self.client.get(reverse('workspace-record-edit', args=['overrides', price.pk])), url)
        response = self.client.get(url)
        self.assertContains(response, 'Delete special price?')
        self.assertContains(response, self.customer.customer_name)
        self.assertContains(response, self.product.product_name)
        self.assertContains(response, '$9.75')
        self.assertTrue(m.CustomerPriceOverride.objects.filter(pk=price.pk).exists())
        self.assertFalse(m.MasterDataChange.objects.filter(source='workspace-delete').exists())
        self.assertRedirects(self.client.post(url), reverse('special-prices'))
        self.assertFalse(m.CustomerPriceOverride.objects.filter(pk=price.pk).exists())
        audit = m.MasterDataChange.objects.get(source='workspace-delete')
        self.assertEqual(audit.actor, self.admin_user)
        self.assertEqual(audit.before['override_price'], '9.75')
        self.assertEqual(audit.after, {})
        self.assertEqual(audit.record_key, str(price.pk))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 6)).price, Decimal('21'))
        order.refresh_from_db()
        self.assertEqual(order.total, Decimal('19.50'))
        self.assertEqual(order.items.get().unit_price, Decimal('9.75'))
        self.assertEqual(self.client.post(url).status_code, 404)
        self.assertEqual(m.MasterDataChange.objects.filter(source='workspace-delete').count(), 1)

    def test_special_price_delete_permission_and_csrf(self):
        price = m.CustomerPriceOverride.objects.create(customer=self.customer, product=self.product,
            override_price='9.75', effective_date=date(2026, 1, 1), active=False)
        url = reverse('special-price-delete', args=[price.pk])
        self.assertTrue(Group.objects.get(name='BAKERY OFFICE').permissions.filter(
            codename='delete_customerpriceoverride').exists())
        denied = get_user_model().objects.create_user('office-no-delete')
        denied.user_permissions.add(Permission.objects.get(codename='office_access'),
            Permission.objects.get(codename='view_customerpriceoverride'))
        restricted = Client()
        restricted.force_login(denied)
        self.assertEqual(restricted.get(url).status_code, 403)
        self.assertEqual(restricted.post(url).status_code, 403)
        self.assertNotContains(restricted.get(reverse('special-prices')), url)
        customer = Client()
        customer.force_login(self.user)
        self.assertEqual(customer.post(url).status_code, 403)
        self.assertEqual(Client().get(url).status_code, 302)
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(self.admin_user)
        secure.get(url)
        self.assertEqual(secure.post(url).status_code, 403)
        self.assertTrue(m.CustomerPriceOverride.objects.filter(pk=price.pk).exists())
        token = secure.cookies['csrftoken'].value
        self.assertRedirects(secure.post(url, {'csrfmiddlewaretoken': token}), reverse('special-prices'))

    def test_special_price_can_be_created_and_edited_for_full_catalog_customer(self):
        from bakery.services.catalog import catalog_links
        self.link.delete()
        before = {link.product_id for link in catalog_links(self.customer)}
        data = {'customer': self.customer.pk, 'product': self.product.pk,
            'override_price': '9.75', 'effective_date': '2026-10-03', 'active': 'on'}
        response = self.client.post(reverse('workspace-record-new', args=['overrides']), data)
        self.assertRedirects(response, reverse('special-prices'))
        saved = m.CustomerPriceOverride.objects.get(customer=self.customer)
        self.assertEqual(saved.created_by, self.admin_user)
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('9.75'))
        data['override_price'] = '8.50'
        response = self.client.post(reverse('workspace-record-edit', args=['overrides', saved.pk]), data)
        self.assertRedirects(response, reverse('special-prices'))
        self.assertEqual(resolve_price(self.customer, self.product, date(2026, 10, 3)).price, Decimal('8.50'))
        self.assertFalse(m.CustomerProduct.objects.filter(customer=self.customer).exists())
        self.assertEqual({link.product_id for link in catalog_links(self.customer)}, before)

    def test_special_price_does_not_reactivate_excluded_or_inactive_products(self):
        self.link.active = False
        self.link.save()
        data = {'customer': self.customer.pk, 'product': self.product.pk,
            'override_price': '9.75', 'effective_date': '2026-10-03', 'active': 'on'}
        self.assertContains(self.client.post(reverse('workspace-record-new', args=['overrides']), data),
            'enable the product there first')
        self.link.delete()
        self.product.active = False
        self.product.save()
        self.assertContains(self.client.post(reverse('workspace-record-new', args=['overrides']), data),
            'Choose an active product')
        self.assertFalse(m.CustomerPriceOverride.objects.exists())

    def test_user_creation_hashes_password_and_does_not_audit_it(self):
        role = Group.objects.get(name='CUSTOMER')
        password = 'Test-only-created-user-338!'
        response = self.client.post(reverse('workspace-user-create'), {
            'username': 'created-customer', 'role': role.pk, 'password1': password, 'password2': password})
        user = get_user_model().objects.get(username='created-customer')
        self.assertRedirects(response, reverse('workspace-user-edit', args=[user.pk]))
        self.assertTrue(user.check_password(password))
        self.assertFalse(user.is_staff)
        self.assertEqual(list(user.groups.values_list('name', flat=True)), ['CUSTOMER'])
        event = m.MasterDataChange.objects.get(source='workspace:user-created')
        self.assertNotIn(password, json.dumps(event.after))
        self.assertNotIn(user.password, json.dumps(event.after))

    def test_customer_and_production_cannot_reach_custom_admin_or_mutate_prices(self):
        self.client.force_login(self.user)
        paths = [reverse('administration'), reverse('workspace-records', args=['overrides']),
                 reverse('workspace-record-new', args=['customers']), reverse('product-detail', args=[self.product.pk]),
                 reverse('workspace-users'), reverse('workspace-user-create'), reverse('workspace-audit')]
        for path in paths:
            self.assertEqual(self.client.get(path).status_code, 403)
        self.assertEqual(self.client.post(reverse('product-prices', args=[self.product.pk]), self.price_post(date(2026, 10, 3))).status_code, 403)
        self.user.groups.add(Group.objects.get(name='PRODUCTION'))
        self.assertEqual(self.client.get(reverse('administration')).status_code, 403)

    def test_office_without_staff_flag_can_use_custom_admin(self):
        self.user.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        self.client.force_login(self.user)
        self.assertFalse(self.user.is_staff)
        self.assertEqual(self.client.get(reverse('administration')).status_code, 200)
        self.assertEqual(self.client.get(reverse('workspace-record-new', args=['products'])).status_code, 200)
        self.assertEqual(self.client.get(reverse('workspace-user-create')).status_code, 403)

    def test_user_role_and_activation_can_be_edited_without_password_change(self):
        original_password = self.user.password
        role = Group.objects.get(name='PRODUCTION')
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), {
            'username': self.user.username, 'role': role.pk, 'first_name': 'Production operator'})
        self.assertRedirects(response, reverse('workspace-users'))
        self.user.refresh_from_db()
        self.assertFalse(self.user.is_active)
        self.assertEqual(self.user.password, original_password)
        self.assertEqual(list(self.user.groups.values_list('name', flat=True)), ['PRODUCTION'])

    def test_administrator_cannot_deactivate_own_login(self):
        role = Group.objects.get(name='ADMINISTRATOR')
        response = self.client.post(reverse('workspace-user-edit', args=[self.admin_user.pk]), {
            'username': self.admin_user.username, 'role': role.pk})
        self.assertContains(response, 'Keep your own administrator account active')
        self.admin_user.refresh_from_db()
        self.assertTrue(self.admin_user.is_active)

    def test_customer_cannot_edit_login_role(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), {
            'role': Group.objects.get(name='ADMINISTRATOR').pk, 'is_active': 'on'}).status_code, 403)

    def access_post(self, rows, include_login=True):
        data = {'access-TOTAL_FORMS': len(rows), 'access-INITIAL_FORMS': 1,
                'access-MIN_NUM_FORMS': 0, 'access-MAX_NUM_FORMS': 200}
        for index, row in enumerate(rows):
            for key, value in row.items():
                data[f'access-{index}-{key}'] = value
        if include_login:
            data.update(username=self.user.username, first_name='Updated person', is_active='on',
                        role=Group.objects.get(name='CUSTOMER').pk)
        return data

    def test_login_and_access_save_together_and_restrict_to_selected_stores(self):
        assignment = self.user.customer_access.get()
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), self.access_post([
            {'id': assignment.pk, 'customer': self.other.pk, 'location': self.second_location.pk, 'active': 'on'},
            {'location': self.other_location.pk, 'active': 'on'},
        ]))
        self.assertRedirects(response, reverse('workspace-users'))
        self.user.refresh_from_db()
        self.assertEqual(self.user.first_name, 'Updated person')
        assignment.refresh_from_db()
        self.assertEqual(assignment.customer_id, self.customer.pk)
        from bakery.services.access import locations_for
        self.assertEqual(set(locations_for(self.user).values_list('pk', flat=True)),
                         {self.second_location.pk, self.other_location.pk})
        self.assertEqual(m.MasterDataChange.objects.filter(source='workspace:login-access').count(), 2)

    def test_invalid_store_assignment_does_not_save_login_or_other_assignments(self):
        assignment = self.user.customer_access.get()
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), self.access_post([
            {'id': assignment.pk, 'customer': self.other.pk, 'location': self.second_location.pk, 'active': 'on'},
            {'location': 999999, 'active': 'on'},
        ]))
        self.assertContains(response, 'Select a valid choice.')
        self.user.refresh_from_db()
        assignment.refresh_from_db()
        self.assertEqual(self.user.first_name, '')
        self.assertEqual(assignment.location_id, self.location.pk)
        self.assertEqual(self.user.customer_access.count(), 1)
        self.assertFalse(m.MasterDataChange.objects.exists())

    def test_access_form_rejects_duplicate_store_assignments_and_foreign_ids(self):
        assignment = self.user.customer_access.get()
        row = {'id': assignment.pk, 'customer': self.customer.pk, 'location': self.location.pk, 'active': 'on'}
        duplicate = {'location': self.other_location.pk, 'active': 'on'}
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]),
                                    self.access_post([row, duplicate, duplicate]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['access_forms'].non_form_errors())
        foreign = m.CustomerAccess.objects.create(user=self.admin_user, customer=self.other, location=self.other_location)
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]),
                                    self.access_post([{**row, 'id': foreign.pk}]))
        self.assertContains(response, 'This assignment belongs to another login.')
        foreign.refresh_from_db()
        self.assertEqual(foreign.user_id, self.admin_user.pk)
        self.assertEqual(foreign.customer_id, self.other.pk)
        self.assertEqual(self.user.customer_access.count(), 1)

    def test_office_can_manage_access_without_changing_login_or_role(self):
        office = get_user_model().objects.create_user('access-office', password='Test-office-only-338!')
        office.groups.add(Group.objects.get(name='BAKERY OFFICE'))
        self.client.force_login(office)
        self.assertContains(self.client.get(reverse('workspace-users')), self.user.username)
        response = self.client.get(reverse('workspace-user-edit', args=[self.user.pk]))
        self.assertNotContains(response, 'name="role"')
        assignment = self.user.customer_access.get()
        rows = [{'id': assignment.pk, 'location': self.second_location.pk, 'active': 'on'}]
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), self.access_post(rows))
        self.assertEqual(response.status_code, 403)
        response = self.client.post(reverse('workspace-user-edit', args=[self.user.pk]), self.access_post(rows, False))
        self.assertRedirects(response, reverse('workspace-users'))
        assignment.refresh_from_db()
        self.user.refresh_from_db()
        self.assertEqual(assignment.location_id, self.second_location.pk)
        self.assertEqual(self.user.first_name, '')
        self.assertFalse(self.user.is_superuser)

    def test_old_customer_access_list_opens_combined_directory(self):
        response = self.client.get(reverse('workspace-records', args=['access']), {'customer': self.customer.pk})
        self.assertRedirects(response, reverse('workspace-users') + f'?customer={self.customer.pk}')
        response = self.client.get(reverse('administration'))
        self.assertContains(response, '<h2>Logins &amp; access</h2>', count=1)
        self.assertNotContains(response, '<h2>Customer access</h2>')

    def test_workspace_edits_require_model_permission(self):
        from django.contrib.auth.models import Permission
        self.user.user_permissions.add(Permission.objects.get(codename='office_access'))
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('workspace-record-new', args=['products'])).status_code, 403)
        self.assertEqual(self.client.post(reverse('product-prices', args=[self.product.pk]), self.price_post(date(2026, 10, 3))).status_code, 403)
